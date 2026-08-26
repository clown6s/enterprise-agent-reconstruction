import io
from docx import Document
from app.ingestion.domain.models import DocumentSource
from app.ingestion.processing.parsers.docx import DocxParser

class FakeImageOcr:
    def __init__(self): self.calls = []
    def extract(self, image, *, filename, index): self.calls.append(index); return f"图{index}文字"

def make_doc():
    doc = Document(); doc.add_heading("一级", level=1); doc.add_paragraph("正文")
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "名称"; table.cell(0, 1).text = "值"
    doc.add_page_break(); doc.add_heading("二级", level=2); doc.add_paragraph("第二页")
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

def test_docx_heading_table_and_page_break():
    result = DocxParser().parse(DocumentSource("a.docx", make_doc()))
    assert "## 一级" in result.text and "### 二级" in result.text
    assert "名称 | 值" in result.text and "[[page:2]]" in result.text
    assert any(block.kind == "table" for block in result.layout_blocks)

def test_docx_image_extractor_is_injected():
    ocr = FakeImageOcr(); result = DocxParser(ocr).parse(DocumentSource("a.docx", make_doc()))
    assert result.parser == "docx" and not ocr.calls
